"""Merge ABDC (2025 + history), JCR 2025 and SCImago 2025 into one compact dataset."""
import csv, json, re, sys
from collections import Counter, defaultdict
import openpyxl
from tqdm import tqdm

def nissn(x):
    if x is None: return ''
    s = re.sub(r'[^0-9Xx]', '', str(x)).upper()
    return s if len(s) == 8 and s != '00000000' else ''

def num(x):
    try:
        if x is None or str(x).strip() in ('', 'N/A', 'n/a', '-'): return None
        return float(str(x).replace(',', '.').strip())
    except: return None

def clean(x):
    return re.sub(r'\s+', ' ', str(x or '')).strip()

FOR = {'3501':'Accounting, auditing and accountability','3502':'Banking, finance and investment','3503':'Business systems in context',
'3504':'Commercial services','3505':'Human resources and industrial relations','3506':'Marketing',
'3507':'Strategy, management and organisational behaviour','3508':'Tourism','3509':'Transportation, logistics and supply chains',
'3599':'Other commerce, management, tourism and services','3801':'Applied economics','3802':'Econometrics','3803':'Economic theory',
'3899':'Other economics','4609':'Information systems','4801':'Commercial law','4905':'Statistics'}

# ---------- ABDC ----------
wb = openpyxl.load_workbook('abdc.xlsx', read_only=True)
abdc = {}   # key -> dict
def abdc_sheet(name, title_col, pub_col, issn_col, eissn_col, for_col, rate_col, start):
    rows = []
    for i, r in enumerate(wb[name].iter_rows(values_only=True)):
        if i < start or not r or r[title_col] is None: continue
        t = clean(r[title_col])
        if t.lower() in ('journal title', 'journal name'): continue
        rating = clean(r[rate_col]).upper()
        if rating not in ('A*', 'A', 'B', 'C'): continue
        rows.append(dict(t=t, pub=clean(r[pub_col]) if pub_col is not None else '', i1=nissn(r[issn_col]), i2=nissn(r[eissn_col]),
                         f=str(r[for_col]).strip()[:4] if r[for_col] is not None else '', r=rating))
    return rows
a25 = abdc_sheet('2025 JQL', 1, 2, 3, 4, 6, 7, 7)
a22 = abdc_sheet('2022 JQL', 0, 1, 2, 3, 5, 6, 8)
a19 = abdc_sheet('2019 JQL', 0, 1, 2, 3, 5, 6, 8)
a16 = abdc_sheet('2016 JQL', 0, 1, 2, 3, 5, 6, 7)
print('ABDC rows', len(a25), len(a22), len(a19), len(a16))
hist = {}
for yr, rows in (('22', a22), ('19', a19), ('16', a16)):
    for r in rows:
        for k in (r['i1'], r['i2']):
            if k: hist.setdefault(k, {})[yr] = r['r']

# ---------- JCR ----------
jws = openpyxl.load_workbook('jcr.xlsx', read_only=True).worksheets[0]
jcr = []
seen = set()
for i, r in enumerate(jws.iter_rows(values_only=True)):
    if i == 0: continue
    key = (clean(r[1]).lower(), nissn(r[5]), nissn(r[6]))
    if key in seen: continue
    seen.add(key)
    jcr.append(dict(t=clean(r[1]), ab=clean(r[3]), pub=clean(r[4]), i1=nissn(r[5]), i2=nissn(r[6]), cites=num(r[7]),
                    cit=num(r[9]), hl=clean(r[10]), jif=num(r[12]), jif5=num(r[13]), jifns=num(r[14]), jci=num(r[15]),
                    q=clean(r[16]) if clean(r[16]) in ('Q1','Q2','Q3','Q4') else '', rk=clean(r[17])))
print('JCR unique', len(jcr))

# ---------- SCImago ----------
sjr = []
with open('sjr.csv', encoding='utf-8') as fh:
    rd = csv.reader(fh, delimiter=';')
    hdr = next(rd)
    H = {h: k for k, h in enumerate(hdr)}
    for r in rd:
        if len(r) < len(hdr): continue
        if r[H['Type']] not in ('journal', 'trade journal'): continue
        issns = [nissn(x) for x in r[H['Issn']].split(',')]
        issns = [x for x in issns if x]
        cats = []
        for c in r[H['Categories']].split(';'):
            c = c.strip()
            m = re.match(r'(.*?)\s*\((Q[1-4])\)$', c)
            if m: cats.append((m.group(1).strip(), m.group(2)))
            elif c: cats.append((c, ''))
        sjr.append(dict(id=r[H['Sourceid']], t=clean(r[H['Title']]), iss=issns, pub=clean(r[H['Publisher']]),
                        oa=r[H['Open Access']] == 'Yes', dia=r[H['Open Access Diamond']] == 'Yes', sjr=num(r[H['SJR']]),
                        q=r[H['SJR Best Quartile']] if r[H['SJR Best Quartile']] in ('Q1','Q2','Q3','Q4') else '',
                        h=num(r[H['H index']]), cpd=num(r[H['Citations / Doc. (2years)']]), docs=num(r[H['Total Docs. (2025)']]),
                        country=clean(r[H['Country']]), cov=clean(r[H['Coverage']]), cats=cats,
                        areas=[a.strip() for a in r[H['Areas']].split(';') if a.strip()]))
print('SJR journals', len(sjr))

# ---------- union by ISSN (union-find) ----------
recs = []  # (src, idx, issns)
for k, r in enumerate(a25): recs.append(('a', k, [x for x in (r['i1'], r['i2']) if x]))
for k, r in enumerate(jcr): recs.append(('j', k, [x for x in (r['i1'], r['i2']) if x]))
for k, r in enumerate(sjr): recs.append(('s', k, r['iss']))
parent = list(range(len(recs)))
def find(x):
    while parent[x] != x:
        parent[x] = parent[parent[x]]; x = parent[x]
    return x
owner = {}
for n, (_, _, iss) in enumerate(recs):
    for s in iss:
        if s in owner: parent[find(n)] = find(owner[s])
        else: owner[s] = n
groups = defaultdict(list)
for n in range(len(recs)): groups[find(n)].append(n)

# ---------- publisher grouping ----------
PUBRULES = [
 ('Elsevier', r'elsevier|pergamon|academic press|cell press|north[- ]holland|mosby|saunders|churchill'),
 ('Springer Nature', r'springer|nature|palgrave|bmc|biomed central|birkh|humana'),
 ('Wiley', r'wiley|blackwell|hindawi|interscience'),
 ('Taylor & Francis', r'taylor|francis|routledge|crc press|informa|psychology press|dove medical'),
 ('SAGE', r'\bsage\b'), ('Emerald', r'emerald'), ('Oxford University Press', r'oxford univ|oup\b|oxford academic'),
 ('Cambridge University Press', r'cambridge univ'), ('IEEE', r'\bieee\b|institute of electrical'), ('ACM', r'\bacm\b|association for computing'),
 ('MDPI', r'mdpi'), ('Frontiers', r'frontiers'), ('De Gruyter', r'gruyter'), ('Inderscience', r'inderscience'), ('IGI Global', r'igi global'),
 ('INFORMS', r'informs|institute for operations research'), ('Academy of Management', r'academy of management'),
 ('American Marketing Association', r'american marketing'), ('American Psychological Association', r'american psychological'),
 ('American Economic Association', r'american economic assoc'), ('University of Chicago Press', r'univ(ersity)? of chicago|u chicago'),
 ('MIT Press', r'mit press'), ('World Scientific', r'world scientific'), ('Wolters Kluwer', r'wolters|lippincott|kluwer'),
 ('IOP Publishing', r'\biop\b|institute of physics'), ('Royal Society of Chemistry', r'royal soc(iety)? (of )?chem'),
 ('American Chemical Society', r'american chemical'), ('Brill', r'\bbrill\b'), ('Thieme', r'thieme'), ('Karger', r'karger'),
 ('PLOS', r'\bplos\b|public library of science'), ('Bentham', r'bentham'), ('Edward Elgar', r'edward elgar'), ('Intellect', r'^intellect'),
 ('Annual Reviews', r'annual reviews'), ('AIP Publishing', r'american institute of physics|\baip\b'), ('APS', r'american physical'),
 ('Duke University Press', r'duke univ'), ('Johns Hopkins University Press', r'johns hopkins'), ('Mary Ann Liebert', r'liebert'),
]
def pubgroup(p):
    pl = p.lower()
    for name, rx in PUBRULES:
        if re.search(rx, pl): return name
    return ''
def titlecase_pub(p):
    if p.isupper(): return ' '.join(w.capitalize() if len(w) > 3 or w in ('INC','LTD') else w for w in p.split()).replace('Inc','Inc').replace(' Of ', ' of ')
    return p

def title_pretty(t):
    if t.isupper() and len(t) > 4:
        small = {'and','of','the','in','for','on','a','an','to','at','by','with','&','de','la','et','des','du','und','fur'}
        ws = t.lower().split(' ')
        out = []
        for i, w in enumerate(ws):
            if i > 0 and w in small: out.append(w)
            elif re.fullmatch(r'(ieee|acm|ii|iii|iv|bmc|bmj|jama|plos|aids|hiv|rna|dna|siam|mis|ssrn|uk|usa|us|eu|asean|it|ai|ieee/acm|cr|ca)', w): out.append(w.upper())
            else: out.append('-'.join(x[:1].upper() + x[1:] for x in w.split('-')))
        return ' '.join(out)
    return t

# ---------- build output ----------
catDict, areaDict, pubDict, grpDict, countryDict = {}, {}, {}, {}, {}
def idx(d, v):
    if v not in d: d[v] = len(d)
    return d[v]

journals = []
stats = Counter()
for g, members in tqdm(groups.items(), desc='merging', file=sys.stdout):
    A = [a25[recs[n][1]] for n in members if recs[n][0] == 'a']
    J = [jcr[recs[n][1]] for n in members if recs[n][0] == 'j']
    S = [sjr[recs[n][1]] for n in members if recs[n][0] == 's']
    a = A[0] if A else None
    j = max(J, key=lambda x: x['jif'] or -1) if J else None
    s = max(S, key=lambda x: x['sjr'] or -1) if S else None
    issns = []
    for x in A: issns += [x['i1'], x['i2']]
    for x in J: issns += [x['i1'], x['i2']]
    for x in S: issns += x['iss']
    issns = list(dict.fromkeys([x for x in issns if x]))
    title = (a['t'] if a else (s['t'] if s else title_pretty(j['t'])))
    pub_raw = (a['pub'] if a and a['pub'] else (s['pub'] if s and s['pub'] else (titlecase_pub(j['pub']) if j else '')))
    grp = pubgroup(' '.join([a['pub'] if a else '', s['pub'] if s else '', j['pub'] if j else '']))
    o = {'t': title}
    if j and j['ab']: o['ab'] = j['ab']
    o['is'] = [f'{x[:4]}-{x[4:]}' for x in issns]
    if pub_raw: o['p'] = idx(pubDict, pub_raw)
    if grp: o['g'] = idx(grpDict, grp)
    if a:
        o['ab25'] = a['r']; o['for'] = a['f']
        h = {}
        for k in issns:
            if k in hist: h.update(hist[k])
        if h: o['ah'] = h
        stats['abdc'] += 1
    if j:
        for k in ('jif','jif5','jifns','jci','cites','cit'):
            if j[k] is not None: o[k] = j[k]
        if j['q']: o['jq'] = j['q']
        if j['rk']: o['jrk'] = j['rk']
        stats['jcr'] += 1
    if s:
        o['sid'] = s['id']
        if s['sjr'] is not None: o['sjr'] = s['sjr']
        if s['q']: o['sq'] = s['q']
        if s['h'] is not None: o['h'] = int(s['h'])
        if s['cpd'] is not None: o['cpd'] = s['cpd']
        if s['docs'] is not None: o['docs'] = int(s['docs'])
        if s['oa']: o['oa'] = 1
        if s['dia']: o['dia'] = 1
        if s['country']: o['c'] = idx(countryDict, s['country'])
        if s['cov']: o['cov'] = s['cov']
        o['cat'] = [[idx(catDict, c), int(q[1]) if q else 0] for c, q in s['cats']]
        o['ar'] = [idx(areaDict, x) for x in s['areas']]
        stats['sjr'] += 1
    journals.append(o)

journals.sort(key=lambda o: o['t'].lower())
print(stats, 'total', len(journals))
meta = dict(cats=list(catDict), areas=list(areaDict), pubs=list(pubDict), groups=list(grpDict), countries=list(countryDict), for_=FOR,
            counts=dict(total=len(journals), abdc=stats['abdc'], jcr=stats['jcr'], sjr=stats['sjr']),
            sources=dict(abdc='ABDC Journal Quality List 2025 (current to Sept 2026)', jcr='Clarivate JCR 2025 release (JIF 2024)',
                         sjr='SCImago Journal Rank 2025 (scimagojr.com, downloaded 27 Oct 2026)'))
json.dump(dict(meta=meta, j=journals), open('merged.json', 'w'), separators=(',', ':'))
# list of Elsevier ISSNs for metric harvesting
els = [o['is'][0] for o in journals if o.get('g') is not None and meta['groups'][o['g']] == 'Elsevier' and (o.get('ab25') or o.get('jif') is not None or o.get('sq'))]
json.dump(els, open('elsevier_issns.json', 'w'))
for name in ('Springer Nature', 'Taylor & Francis', 'Wiley', 'SAGE', 'Emerald'):
    print(name, sum(1 for o in journals if o.get('g') is not None and meta['groups'][o['g']] == name), 'abdc:',
          sum(1 for o in journals if o.get('ab25') and o.get('g') is not None and meta['groups'][o['g']] == name))
print('Elsevier targets', len(els))
